"""
Servicio especializado para generación, formato y procesamiento de Libros Contables (Ventas y Compras).
Maneja construcción de hojas de cálculo ejecutivas OpenPyXL, estilos corporativos e importación masiva.
"""
import io
import uuid
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
from fastapi import HTTPException, Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.models.models import (
    AccountReceivable,
    AccountPayable,
    Client,
    Project,
    Expense,
    ExpenseCategory,
    AuditLog
)

class FinancialExcelService:

    @staticmethod
    def get_libro_ventas_data(
        month: Optional[int] = None,
        year: Optional[int] = None,
        filter_type: Optional[str] = None,
        db: Session = None
    ):
        """
        Retorna los datos del Libro de Ventas en formato oficial SENIAT para renderizado y PDF.
        filter_type: 'all', 'fiscales', 'no_fiscales'
        """
        query = db.query(AccountReceivable).options(
            joinedload(AccountReceivable.client),
            joinedload(AccountReceivable.project)
        )
        if year:
            query = query.filter(func.extract('year', AccountReceivable.issue_date) == year)
        if month:
            query = query.filter(func.extract('month', AccountReceivable.issue_date) == month)
    
        if filter_type in ("fiscales", "con_iva"):
            query = query.filter(
                AccountReceivable.tax_amount_usd > 0.001,
                ~AccountReceivable.invoice_number.ilike("NE-%"),
                ~AccountReceivable.invoice_number.ilike("NDE-%"),
                ~AccountReceivable.invoice_number.ilike("REC-%"),
                ~AccountReceivable.invoice_number.ilike("VAL-%")
            )
        elif filter_type == "sin_iva":
            query = query.filter(
                (AccountReceivable.tax_amount_usd == 0) | (AccountReceivable.tax_amount_usd == None) | (AccountReceivable.tax_amount_usd <= 0.001),
                ~AccountReceivable.invoice_number.ilike("NE-%"),
                ~AccountReceivable.invoice_number.ilike("NDE-%"),
                ~AccountReceivable.invoice_number.ilike("REC-%"),
                ~AccountReceivable.invoice_number.ilike("VAL-%")
            )
        elif filter_type == "no_fiscales":
            query = query.filter(
                (AccountReceivable.invoice_number.ilike("NE-%")) |
                (AccountReceivable.invoice_number.ilike("NDE-%")) |
                (AccountReceivable.invoice_number.ilike("REC-%")) |
                (AccountReceivable.invoice_number.ilike("VAL-%"))
            )
    
        receivables = query.order_by(AccountReceivable.issue_date.asc(), AccountReceivable.id.asc()).all()
    
        items = []
        tot_ventas_bs = 0.0
        tot_base_bs = 0.0
        tot_iva_bs = 0.0
        tot_ret_iva_bs = 0.0
        tot_ret_islr_bs = 0.0
    
        for idx, r in enumerate(receivables, start=1):
            rate = r.exchange_rate or 859.06
            total_bs = round(r.amount_usd * rate, 2)
            tax_usd = r.tax_amount_usd or 0.0
            base_usd = r.taxable_base_usd or 0.0
    
            if tax_usd <= 0.001:
                iva_bs = 0.0
                base_bs = 0.0
            else:
                iva_bs = round(tax_usd * rate, 2)
                base_bs = round(base_usd * rate, 2) if base_usd > 0 else round(total_bs / 1.16, 2)
    
            ret_iva_bs = round((r.tax_withholding_usd or 0.0) * rate, 2)
            ret_islr_bs = round((r.islr_withholding_usd or 0.0) * rate, 2)
    
            tot_ventas_bs += total_bs
            tot_base_bs += base_bs
            tot_iva_bs += iva_bs
            tot_ret_iva_bs += ret_iva_bs
            tot_ret_islr_bs += ret_islr_bs
    
            items.append({
                "operacion": idx,
                "id": r.id,
                "fecha": r.issue_date.strftime("%d/%m/%Y") if r.issue_date else "-",
                "dia": r.issue_date.day if r.issue_date else 1,
                "factura": r.invoice_number,
                "control": r.invoice_number or f"00-{r.id:06d}",
                "tipo_transaccion": "01",
                "cliente": r.client.name if r.client else "Cliente General",
                "rif": r.client.rif if r.client else "J-00000000-0",
                "cbt_retencion": r.notes or "-",
                "ret_iva_bs": ret_iva_bs,
                "ret_islr_bs": ret_islr_bs,
                "total_ventas_bs": total_bs,
                "base_imponible_bs": base_bs,
                "alicuota_pct": 16.0 if tax_usd > 0.001 else 0.0,
                "iva_bs": iva_bs,
                "total_ventas_usd": r.amount_usd,
                "exchange_rate": rate
            })
    
        return {
            "success": True,
            "company": {
                "name": "METALMECANICA DALOR, C.A.",
                "rif": "J-31601195-0",
                "address": "AV. CAMARA DE LAS INDUSTRIAS ZONA INDUSTRIAL EL TIGRE GALPON N° 10 GUACARA EDO. CARABOBO",
                "period": f"{month or 'Todos'}/{year or '2026'}"
            },
            "totals": {
                "total_ventas_bs": round(tot_ventas_bs, 2),
                "base_imponible_bs": round(tot_base_bs, 2),
                "iva_bs": round(tot_iva_bs, 2),
                "ret_iva_bs": round(tot_ret_iva_bs, 2),
                "ret_islr_bs": round(tot_ret_islr_bs, 2)
            },
            "items": items
        }

    @staticmethod
    def export_libro_ventas_excel(
        month: Optional[int] = None,
        year: Optional[int] = None,
        filter_type: Optional[str] = None,
        db: Session = None
    ):
        """
        Exporta el Libro de Ventas en formato claro, armónico y profesional de Excel.
        """
        data = FinancialExcelService.get_libro_ventas_data(month=month, year=year, filter_type=filter_type, db=db)
        items = data["items"]
        totals = data["totals"]
    
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Libro de Ventas"
        ws.views.sheetView[0].showGridLines = True
    
        font_company = Font(name="Calibri", size=13, bold=True, color="002B49")
        font_sub = Font(name="Calibri", size=9, color="555555")
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_regular = Font(name="Calibri", size=9)
        font_title = Font(name="Calibri", size=14, bold=True, color="002B49")
        font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    
        fill_header = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
        fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        fill_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_totals = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
        border_cell = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        border_totals = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='002B49'),
            bottom=Side(style='double', color='002B49')
        )
    
        # Encabezado corporativo
        ws['A1'] = "METALMECÁNICA DALOR, C.A."
        ws['A1'].font = font_company
        ws['A2'] = "AV. CÁMARA DE LAS INDUSTRIAS, ZONA INDUSTRIAL EL TIGRE, GUACARA, EDO. CARABOBO"
        ws['A2'].font = font_sub
        ws['A3'] = "RIF: J-31601195-0"
        ws['A3'].font = font_bold
        ws['E3'] = "LIBRO DE VENTAS"
        ws['E3'].font = font_title
        ws['L3'] = f"PERÍODO: {month or 'TODOS'}/{year or '2026'}"
        ws['L3'].font = font_bold
    
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 15
        ws.row_dimensions[3].height = 24
        ws.row_dimensions[4].height = 10
    
        headers = [
            "Fecha",
            "N° Factura",
            "N° Control",
            "Tipo",
            "Cliente / Razón Social",
            "RIF",
            "N° Cbt. Ret.",
            "Total Ventas (Bs.)",
            "Base Imponible (Bs.)",
            "% Alícuota",
            "IVA Débito (Bs.)",
            "Ret. IVA (Bs.)",
            "Ret. ISLR (Bs.)"
        ]
    
        ws.row_dimensions[5].height = 28
        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=5, column=c_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell
    
        row_num = 6
        for idx, item in enumerate(items):
            alicuota_str = f"{int(item['alicuota_pct'])}%" if item.get('alicuota_pct') else ("16%" if item.get('iva_bs', 0) > 0 else "0%")
            cbt_ret = item.get("cbt_retencion")
            if cbt_ret == "-":
                cbt_ret = ""
            ws.append([
                item.get("fecha", "-"),
                item.get("factura", "-"),
                item.get("control", "-"),
                "01",
                item.get("cliente", "-"),
                item.get("rif", "-"),
                cbt_ret,
                item.get("total_ventas_bs", 0.0),
                item.get("base_imponible_bs", 0.0),
                alicuota_str,
                item.get("iva_bs", 0.0),
                item.get("ret_iva_bs", 0.0),
                item.get("ret_islr_bs", 0.0)
            ])
            ws.row_dimensions[row_num].height = 20
            fill_curr = fill_odd if idx % 2 == 1 else fill_even
    
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row_num, c_idx)
                c.font = font_regular
                c.fill = fill_curr
                c.border = border_cell
                if c_idx in [8, 9, 11, 12, 13]:
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 5:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center")
            row_num += 1
    
        # Totales
        ws.append([
            "", "", "", "", "TOTALES GENERALES", "", "",
            totals.get("total_ventas_bs", 0.0),
            totals.get("base_imponible_bs", 0.0),
            "16%",
            totals.get("iva_bs", 0.0),
            totals.get("ret_iva_bs", 0.0),
            totals.get("ret_islr_bs", 0.0)
        ])
        ws.row_dimensions[row_num].height = 24
        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row_num, c_idx)
            c.font = Font(name="Calibri", size=9, bold=True, color="002B49")
            c.fill = fill_totals
            c.border = border_totals
            if c_idx in [8, 9, 11, 12, 13]:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 5:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
    
        # Auto ajuste de anchos
        for c_idx in range(1, len(headers) + 1):
            col_letter = openpyxl.utils.get_column_letter(c_idx)
            max_len = len(str(headers[c_idx - 1]))
            for r_idx in range(5, row_num + 1):
                val = ws.cell(r_idx, c_idx).value
                if val is not None:
                    val_str = f"{val:,.2f}" if isinstance(val, (int, float)) else str(val)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 45)
    
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
    
        fname = f"Libro_de_Ventas_{year or 2026}_{month or 'Todos'}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={fname}"}
        )
    

    @staticmethod
    def get_libro_compras_data(
        month: Optional[int] = None,
        year: Optional[int] = None,
        filter_type: Optional[str] = None,
        db: Session = None
    ):
        """
        Retorna los datos del Libro de Compras en formato oficial SENIAT (21 columnas) para renderizado y PDF.
        filter_type: 'all', 'fiscales', 'no_fiscales'
        """
        query = db.query(AccountPayable).options(
            joinedload(AccountPayable.project)
        )
        if year:
            query = query.filter(func.extract('year', AccountPayable.issue_date) == year)
        if month:
            query = query.filter(func.extract('month', AccountPayable.issue_date) == month)
    
        if filter_type in ("fiscales", "con_iva"):
            query = query.filter(
                AccountPayable.tax_amount_usd > 0.001,
                AccountPayable.doc_type != "nota_entrega",
                ~AccountPayable.invoice_number.ilike("NE-%"),
                ~AccountPayable.invoice_number.ilike("NDE-%"),
                ~AccountPayable.invoice_number.ilike("REC-%"),
                ~AccountPayable.invoice_number.ilike("VAL-%")
            )
        elif filter_type == "sin_iva":
            query = query.filter(
                (AccountPayable.tax_amount_usd == 0) | (AccountPayable.tax_amount_usd == None) | (AccountPayable.tax_amount_usd <= 0.001),
                AccountPayable.doc_type != "nota_entrega",
                ~AccountPayable.invoice_number.ilike("NE-%"),
                ~AccountPayable.invoice_number.ilike("NDE-%"),
                ~AccountPayable.invoice_number.ilike("REC-%"),
                ~AccountPayable.invoice_number.ilike("VAL-%")
            )
        elif filter_type == "no_fiscales":
            query = query.filter(
                (AccountPayable.doc_type == "nota_entrega") |
                (AccountPayable.invoice_number.ilike("NE-%")) |
                (AccountPayable.invoice_number.ilike("NDE-%")) |
                (AccountPayable.invoice_number.ilike("REC-%")) |
                (AccountPayable.invoice_number.ilike("VAL-%"))
            )
    
        payables = query.order_by(AccountPayable.issue_date.asc(), AccountPayable.id.asc()).all()
    
        items = []
        tot_compras_bs = 0.0
        tot_ret_iva_bs = 0.0
        tot_ret_islr_bs = 0.0
        tot_ret_mun_bs = 0.0
        tot_exento_bs = 0.0
        tot_base_16_bs = 0.0
        tot_iva_16_bs = 0.0
    
        for idx, p in enumerate(payables, start=1):
            rate = p.exchange_rate or 859.06
            total_bs = round(p.amount_usd * rate, 2)
            exento_usd = p.withholding_exempt_usd or 0.0
            tax_usd = p.tax_amount_usd or 0.0
            base_usd = p.taxable_base_usd or 0.0
    
            if tax_usd <= 0.001:
                iva_16_bs = 0.0
                base_16_bs = 0.0
                exento_bs = total_bs
            else:
                iva_16_bs = round(tax_usd * rate, 2)
                base_16_bs = round(base_usd * rate, 2) if base_usd > 0 else round((total_bs - round(exento_usd * rate, 2)) / 1.16, 2)
                exento_bs = round(exento_usd * rate, 2)
    
            ret_iva_bs = round((p.tax_withholding_usd or 0.0) * rate, 2)
            ret_islr_bs = round((p.islr_withholding_usd or 0.0) * rate, 2)
            ret_mun_bs = round((p.municipal_withholding_usd or 0.0) * rate, 2)
    
            tot_compras_bs += total_bs
            tot_ret_iva_bs += ret_iva_bs
            tot_ret_islr_bs += ret_islr_bs
            tot_ret_mun_bs += ret_mun_bs
            tot_exento_bs += exento_bs
            tot_base_16_bs += base_16_bs
            tot_iva_16_bs += iva_16_bs
    
            items.append({
                "operacion": idx,
                "id": p.id,
                "payable_id": p.id,
                "fecha": p.issue_date.strftime("%d/%m/%Y") if p.issue_date else "-",
                "dia": p.issue_date.day if p.issue_date else 1,
                "factura": p.invoice_number,
                "control": p.control_number or f"00-{p.id:06d}",
                "tipo_transaccion": "01",
                "proveedor": p.supplier_name,
                "rif": p.supplier_rif or "J-00000000-0",
                "cbt_retencion": p.withholding_voucher_number or "-",
                "total_compras_bs": total_bs,
                "ret_iva_bs": ret_iva_bs,
                "ret_islr_bs": ret_islr_bs,
                "ret_municipal_bs": ret_mun_bs,
                "exento_bs": exento_bs,
                "base_8_bs": 0.0,
                "alicuota_8_pct": 8.0,
                "iva_8_bs": 0.0,
                "base_12_bs": 0.0,
                "alicuota_12_pct": 12.0,
                "iva_12_bs": 0.0,
                "base_16_bs": base_16_bs,
                "alicuota_16_pct": 16.0 if tax_usd > 0.001 else 0.0,
                "iva_16_bs": iva_16_bs,
                "total_usd": p.amount_usd,
                "exchange_rate": rate,
                "municipal_rate": p.municipal_rate or 0.0,
                "municipal_voucher_number": p.municipal_voucher_number
            })
    
        return {
            "success": True,
            "company": {
                "name": "METALMECANICA DALOR, C.A.",
                "rif": "J-31601195-0",
                "address": "AV. CAMARA DE LAS INDUSTRIAS ZONA INDUSTRIAL EL TIGRE GUACARA CARABOBO",
                "period": f"{month or 'Todos'}/{year or '2026'}"
            },
            "totals": {
                "total_compras_bs": round(tot_compras_bs, 2),
                "ret_iva_bs": round(tot_ret_iva_bs, 2),
                "ret_islr_bs": round(tot_ret_islr_bs, 2),
                "ret_mun_bs": round(tot_ret_mun_bs, 2),
                "exento_bs": round(tot_exento_bs, 2),
                "base_16_bs": round(tot_base_16_bs, 2),
                "iva_16_bs": round(tot_iva_16_bs, 2)
            },
            "items": items
        }

    @staticmethod
    def export_libro_compras_excel(
        month: Optional[int] = None,
        year: Optional[int] = None,
        filter_type: Optional[str] = None,
        db: Session = None
    ):
        """
        Exporta el Libro de Compras en formato claro, armónico y profesional de Excel.
        """
        data = FinancialExcelService.get_libro_compras_data(month=month, year=year, filter_type=filter_type, db=db)
        items = data["items"]
        totals = data["totals"]
    
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Libro de Compras"
        ws.views.sheetView[0].showGridLines = True
    
        font_company = Font(name="Calibri", size=13, bold=True, color="002B49")
        font_sub = Font(name="Calibri", size=9, color="555555")
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_regular = Font(name="Calibri", size=9)
        font_title = Font(name="Calibri", size=14, bold=True, color="002B49")
        font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    
        fill_header = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
        fill_even = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        fill_odd = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_totals = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
        border_cell = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )
        border_totals = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='002B49'),
            bottom=Side(style='double', color='002B49')
        )
    
        # Encabezado corporativo
        ws['A1'] = "METALMECÁNICA DALOR, C.A."
        ws['A1'].font = font_company
        ws['A2'] = "AV. CÁMARA DE LAS INDUSTRIAS, ZONA INDUSTRIAL EL TIGRE, GUACARA, EDO. CARABOBO"
        ws['A2'].font = font_sub
        ws['A3'] = "RIF: J-31601195-0"
        ws['A3'].font = font_bold
        ws['E3'] = "LIBRO DE COMPRAS"
        ws['E3'].font = font_title
        ws['M3'] = f"PERÍODO: {month or 'TODOS'}/{year or '2026'}"
        ws['M3'].font = font_bold
    
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 15
        ws.row_dimensions[3].height = 24
        ws.row_dimensions[4].height = 10
    
        headers = [
            "Fecha",
            "N° Factura",
            "N° Control",
            "Tipo",
            "Proveedor / Razón Social",
            "RIF Proveedor",
            "N° Cbt. Ret.",
            "Total Compras (Bs.)",
            "Compras Exentas (Bs.)",
            "Base Imponible (Bs.)",
            "% Alícuota",
            "IVA Crédito (Bs.)",
            "Ret. IVA (Bs.)",
            "Ret. ISLR (Bs.)",
            "Ret. Municipal (Bs.)"
        ]
    
        ws.row_dimensions[5].height = 28
        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=5, column=c_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell
    
        row_num = 6
        for idx, item in enumerate(items):
            alicuota_str = f"{int(item['alicuota_16_pct'])}%" if item.get('alicuota_16_pct') else ("16%" if item.get('iva_16_bs', 0) > 0 else "0%")
            cbt_ret = item.get("cbt_retencion")
            if cbt_ret == "-":
                cbt_ret = ""
            ws.append([
                item.get("fecha", "-"),
                item.get("factura", "-"),
                item.get("control", "-"),
                "01",
                item.get("proveedor", "-"),
                item.get("rif", "-"),
                cbt_ret,
                item.get("total_compras_bs", 0.0),
                item.get("exento_bs", 0.0),
                item.get("base_16_bs", 0.0),
                alicuota_str,
                item.get("iva_16_bs", 0.0),
                item.get("ret_iva_bs", 0.0),
                item.get("ret_islr_bs", 0.0),
                item.get("ret_municipal_bs", 0.0)
            ])
            ws.row_dimensions[row_num].height = 20
            fill_curr = fill_odd if idx % 2 == 1 else fill_even
    
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row_num, c_idx)
                c.font = font_regular
                c.fill = fill_curr
                c.border = border_cell
                if c_idx in [8, 9, 10, 12, 13, 14, 15]:
                    c.number_format = '#,##0.00'
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 5:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center")
            row_num += 1
    
        # Totales Compras
        ws.append([
            "", "", "", "", "TOTALES GENERALES", "", "",
            totals.get("total_compras_bs", 0.0),
            totals.get("exento_bs", 0.0),
            totals.get("base_16_bs", 0.0),
            "16%",
            totals.get("iva_16_bs", 0.0),
            totals.get("ret_iva_bs", 0.0),
            totals.get("ret_islr_bs", 0.0),
            totals.get("ret_mun_bs", 0.0)
        ])
        ws.row_dimensions[row_num].height = 24
        for c_idx in range(1, len(headers) + 1):
            c = ws.cell(row_num, c_idx)
            c.font = Font(name="Calibri", size=9, bold=True, color="002B49")
            c.fill = fill_totals
            c.border = border_totals
            if c_idx in [8, 9, 10, 12, 13, 14, 15]:
                c.number_format = '#,##0.00'
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 5:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
    
        # Auto ajuste de anchos
        for c_idx in range(1, len(headers) + 1):
            col_letter = openpyxl.utils.get_column_letter(c_idx)
            max_len = len(str(headers[c_idx - 1]))
            for r_idx in range(5, row_num + 1):
                val = ws.cell(r_idx, c_idx).value
                if val is not None:
                    val_str = f"{val:,.2f}" if isinstance(val, (int, float)) else str(val)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 45)
    
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
    
        fname = f"Libro_de_Compras_{year or 2026}_{month or 'Todos'}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={fname}"}
        )
    
    async def import_libros_excel(
        file: Any = None,
        exchange_rate: float = None,
        db: Session = None
    ):
        """
        Importa masivamente facturas y retenciones desde un archivo Excel de Libros de Compras y Ventas.
        """
        contents = await file.read()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo Excel: {str(e)}")
    
        def safe_float(v):
            if v is None: return 0.0
            try:
                return float(str(v).replace(',', '.').strip())
            except:
                return 0.0
    
        imported_ventas = 0
        imported_compras = 0
        linked_municipal = 0
    
        rate = exchange_rate or 859.06
    
        # 1. Procesar Ventas
        if 'Ventas' in wb.sheetnames:
            ws_v = wb['Ventas']
            for r in range(6, ws_v.max_row + 1):
                factura = ws_v.cell(r, 2).value
                cliente_nombre = ws_v.cell(r, 5).value
                if factura and cliente_nombre and str(factura).strip() != 'FACTURA' and 'TOTAL' not in str(cliente_nombre):
                    fact_clean = str(factura).strip()
                    cli_clean = str(cliente_nombre).strip()
                    rif_clean = str(ws_v.cell(r, 9).value or '').strip()
                    control_clean = str(ws_v.cell(r, 3).value or fact_clean).strip()
    
                    # Buscar o crear Cliente
                    client = db.query(Client).filter(Client.name.ilike(cli_clean)).first()
                    if not client:
                        client = Client(name=cli_clean, rif=rif_clean or "J-00000000-0", contact_name="Contacto Comercial")
                        db.add(client)
                        db.flush()
    
                    # Buscar duplicado en AccountReceivable
                    existing_rec = db.query(AccountReceivable).filter(AccountReceivable.invoice_number == fact_clean).first()
                    if not existing_rec:
                        dia_val = ws_v.cell(r, 1).value
                        issue_dt = dia_val if isinstance(dia_val, datetime) else datetime(2026, 9, 9)
    
                        tot_bs = safe_float(ws_v.cell(r, 10).value)
                        base_bs = safe_float(ws_v.cell(r, 11).value)
                        ret_iva_bs = safe_float(ws_v.cell(r, 7).value)
                        ret_islr_bs = safe_float(ws_v.cell(r, 8).value)
    
                        tot_usd = round(tot_bs / rate, 2)
                        base_usd = round(base_bs / rate, 2)
                        tax_usd = round((tot_bs - base_bs) / rate, 2)
                        ret_iva_usd = round(ret_iva_bs / rate, 2)
                        ret_islr_usd = round(ret_islr_bs / rate, 2)
                        net_usd = round(tot_usd - ret_iva_usd - ret_islr_usd, 2)
    
                        new_rec = AccountReceivable(
                            invoice_number=fact_clean,
                            client_id=client.id,
                            description=f"Facturación según Libro de Ventas SENIAT #{fact_clean}",
                            issue_date=issue_dt,
                            due_date=issue_dt + timedelta(days=30),
                            amount_usd=tot_usd,
                            amount_bs=tot_bs,
                            exchange_rate=rate,
                            taxable_base_usd=base_usd,
                            tax_amount_usd=tax_usd,
                            tax_withholding_rate=75.0 if ret_iva_bs > 0 else 0.0,
                            tax_withholding_usd=ret_iva_usd,
                            islr_rate=2.0 if ret_islr_bs > 0 else 0.0,
                            islr_withholding_usd=ret_islr_usd,
                            net_amount_usd=net_usd,
                            paid_amount_usd=round(ret_iva_usd + ret_islr_usd, 2),
                            balance_usd=net_usd,
                            status="pendiente" if net_usd > 0.01 else "cobrado_total",
                            notes=f"Control fiscal: {control_clean}. RIF: {rif_clean}"
                        )
                        db.add(new_rec)
                        imported_ventas += 1
    
        # 2. Procesar Compras
        if 'Compras' in wb.sheetnames:
            ws_c = wb['Compras']
            for r in range(7, ws_c.max_row + 1):
                factura = ws_c.cell(r, 2).value
                prov_nombre = ws_c.cell(r, 5).value
                if factura and prov_nombre and str(factura).strip() != 'NUMERO' and 'TOTAL' not in str(prov_nombre):
                    fact_clean = str(factura).strip()
                    prov_clean = str(prov_nombre).strip()
                    rif_clean = str(ws_c.cell(r, 6).value or '').strip()
                    control_clean = str(ws_c.cell(r, 3).value or fact_clean).strip()
                    cbt_ret = str(ws_c.cell(r, 7).value or '').strip()
    
                    # Buscar duplicado en AccountPayable
                    existing_pay = db.query(AccountPayable).filter(
                        AccountPayable.invoice_number == fact_clean,
                        AccountPayable.supplier_name.ilike(prov_clean)
                    ).first()
                    if not existing_pay:
                        dia_val = ws_c.cell(r, 1).value
                        issue_dt = dia_val if isinstance(dia_val, datetime) else datetime(2026, 9, 1)
    
                        tot_bs = safe_float(ws_c.cell(r, 8).value)
                        ret_iva_bs = safe_float(ws_c.cell(r, 9).value)
                        exento_bs = safe_float(ws_c.cell(r, 10).value)
                        base_16_bs = safe_float(ws_c.cell(r, 17).value)
                        iva_16_bs = safe_float(ws_c.cell(r, 19).value)
    
                        tot_usd = round(tot_bs / rate, 2)
                        exento_usd = round(exento_bs / rate, 2)
                        base_usd = round(base_16_bs / rate, 2) if base_16_bs > 0 else round((tot_bs - exento_bs) / (1.16 * rate), 2)
                        tax_usd = round(iva_16_bs / rate, 2) if iva_16_bs > 0 else round((tot_bs - exento_bs - (base_usd * rate)) / rate, 2)
                        ret_iva_usd = round(ret_iva_bs / rate, 2)
                        net_usd = round(tot_usd - ret_iva_usd, 2)
    
                        new_pay = AccountPayable(
                            invoice_number=fact_clean,
                            control_number=control_clean,
                            supplier_name=prov_clean,
                            supplier_rif=rif_clean or None,
                            doc_type="factura",
                            description=f"Compra/Insumos según Libro de Compras SENIAT #{fact_clean}",
                            issue_date=issue_dt,
                            due_date=issue_dt + timedelta(days=15),
                            amount_usd=tot_usd,
                            amount_bs=tot_bs,
                            exchange_rate=rate,
                            taxable_base_usd=base_usd,
                            tax_amount_usd=tax_usd,
                            withholding_exempt_usd=exento_usd,
                            tax_withholding_rate=75.0 if ret_iva_bs > 0 else 0.0,
                            tax_withholding_usd=ret_iva_usd,
                            withholding_voucher_number=cbt_ret or None,
                            withholding_voucher_date=issue_dt if cbt_ret else None,
                            is_withholding_applied=bool(ret_iva_bs > 0),
                            net_amount_usd=net_usd,
                            paid_amount_usd=ret_iva_usd,
                            balance_usd=net_usd,
                            status="pagado_total" if net_usd <= 0.01 else "pendiente"
                        )
                        db.add(new_pay)
                        imported_compras += 1
    
        # 3. Procesar Retención Municipal si está presente
        if 'Municipal' in wb.sheetnames:
            ws_m = wb['Municipal']
            fact_mun = str(ws_m.cell(33, 4).value or '').strip()
            prov_mun = str(ws_m.cell(26, 5).value or '').strip()
            base_mun_bs = safe_float(ws_m.cell(38, 4).value)
            ret_mun_bs = safe_float(ws_m.cell(39, 8).value)
            tasa_mun_pct = safe_float(ws_m.cell(39, 4).value) * 100
            voucher_mun_raw = str(ws_m.cell(9, 9).value or '00000110').replace('Comprobante  N', '').replace('Comprobante', '').strip()
    
            if fact_mun or prov_mun:
                target_pay = db.query(AccountPayable).filter(
                    (AccountPayable.invoice_number == fact_mun) | (AccountPayable.supplier_name.ilike(f"%{prov_mun[:15]}%"))
                ).first()
                if target_pay:
                    target_pay.municipal_rate = tasa_mun_pct or 3.0
                    target_pay.municipal_withholding_usd = round(ret_mun_bs / rate, 2)
                    target_pay.municipal_voucher_number = voucher_mun_raw
                    target_pay.municipal_voucher_date = datetime(2026, 9, 10)
                    linked_municipal += 1
    
        db.commit()
    
        return {
            "success": True,
            "message": f"¡Importación completada! Se registraron {imported_ventas} ventas en CxC, {imported_compras} compras en CxP y se vincularon {linked_municipal} retenciones municipales.",
            "imported_ventas": imported_ventas,
            "imported_compras": imported_compras,
            "linked_municipal": linked_municipal
        }
    
    
