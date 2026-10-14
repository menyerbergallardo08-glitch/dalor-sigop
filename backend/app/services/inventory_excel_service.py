import io
from typing import Optional, List, Dict, Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.models.models import Material, MaterialMovement, Asset, AuditLog


class InventoryExcelService:
    """
    Servicio integral para la exportación e importación masiva de inventarios
    (Materiales, Herramientas, Maquinarias, Equipos) en formato Excel (.xlsx).
    """

    @staticmethod
    def _create_header_style():
        navy_fill = PatternFill(start_color="002B49", end_color="002B49", fill_type="solid")
        blue_fill = PatternFill(start_color="0072B8", end_color="0072B8", fill_type="solid")
        title_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        return navy_fill, blue_fill, title_font, header_font, thin_border

    # =========================================================================
    # 1. MATERIALES: PLANTILLA & EXPORTACIÓN
    # =========================================================================
    @staticmethod
    def generate_materials_template(db: Session) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Catalogo_Materiales"
        ws.views.sheetView[0].showGridLines = True

        navy_fill, blue_fill, title_font, header_font, thin_border = InventoryExcelService._create_header_style()
        regular_font = Font(name="Calibri", size=10)
        italic_hint = Font(name="Calibri", size=9, italic=True, color="64748B")

        # Título institucional
        ws.merge_cells("A1:H1")
        ws["A1"] = "DALOR C.A. | SIGO-P - REFORMA Y CATÁLOGO DE MATERIALES E INSUMOS"
        ws["A1"].font = title_font
        ws["A1"].fill = navy_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:H2")
        ws["A2"] = "Instrucciones: Puede corregir registros existentes (por CÓDIGO) o agregar nuevos. Si modifica el Stock Actual, el sistema registrará un movimiento de ajuste automático."
        ws["A2"].font = italic_hint
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        headers = [
            ("CODIGO", 14),
            ("NOMBRE_MATERIAL", 38),
            ("CATEGORIA", 24),
            ("UNIDAD_MEDIDA", 16),
            ("STOCK_ACTUAL", 16),
            ("STOCK_MINIMO", 16),
            ("COSTO_UNITARIO_USD", 20),
            ("UBICACION", 28)
        ]

        ws.row_dimensions[3].height = 24
        for col_idx, (col_name, col_width) in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = blue_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        # Precargar materiales existentes de la base de datos
        materials = db.query(Material).filter((Material.is_active == True) | (Material.is_active == None)).order_by(Material.code.asc()).all()

        current_row = 4
        for mat in materials:
            ws.row_dimensions[current_row].height = 19
            row_data = [
                mat.code,
                mat.name,
                mat.category or "Acero Estructural",
                mat.unit_measure or "UND",
                float(mat.stock_quantity or 0.0),
                float(mat.min_stock_alert or 5.0),
                float(mat.unit_cost_usd or 0.0),
                mat.location or "Almacen Central Dalor"
            ]
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in (5, 6, 7):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if col_idx == 7:
                        cell.number_format = '"$"#,##0.00'
                    else:
                        cell.number_format = '#,##0.00'
                elif col_idx in (1, 4):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # =========================================================================
    # 2. MATERIALES: IMPORTACIÓN MASIVA
    # =========================================================================
    @staticmethod
    def import_materials_excel(file_bytes: bytes, db: Session) -> Dict[str, Any]:
        wb = load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        created_count = 0
        updated_count = 0
        adjusted_stock_count = 0
        errors = []

        # Buscar fila de cabecera (debe contener CODIGO y NOMBRE_MATERIAL)
        header_row = 3
        headers_map = {}
        for row_idx in range(1, 6):
            row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip().upper() for c in range(1, 15)]
            if "CODIGO" in row_vals and any("NOMBRE" in v for v in row_vals):
                header_row = row_idx
                for col_idx in range(1, 15):
                    v = str(ws.cell(row=row_idx, column=col_idx).value or "").strip().upper()
                    if v:
                        headers_map[v] = col_idx
                break

        if not headers_map or "CODIGO" not in headers_map:
            raise ValueError("El archivo no tiene el formato esperado. Faltan las columnas CODIGO o NOMBRE_MATERIAL.")

        max_row = ws.max_row
        for r in range(header_row + 1, max_row + 1):
            def get_val(key: str) -> Optional[Any]:
                col = headers_map.get(key)
                if not col:
                    return None
                val = ws.cell(row=r, column=col).value
                return val

            raw_code = str(get_val("CODIGO") or "").strip().upper()
            raw_name = str(get_val("NOMBRE_MATERIAL") or "").strip()

            if not raw_code and not raw_name:
                continue

            if not raw_name:
                errors.append(f"Fila {r}: El registro con código {raw_code} no tiene nombre.")
                continue

            category = str(get_val("CATEGORIA") or "Acero Estructural").strip()
            unit = str(get_val("UNIDAD_MEDIDA") or "UND").strip().upper()
            location = str(get_val("UBICACION") or "Almacen Central Dalor").strip()

            try:
                stock_qty = max(0.0, float(get_val("STOCK_ACTUAL") or 0.0))
            except (ValueError, TypeError):
                stock_qty = 0.0

            try:
                min_stock = max(0.0, float(get_val("STOCK_MINIMO") or 5.0))
            except (ValueError, TypeError):
                min_stock = 5.0

            try:
                unit_cost = max(0.0, float(get_val("COSTO_UNITARIO_USD") or 0.0))
            except (ValueError, TypeError):
                unit_cost = 0.0

            # Buscar por código existente
            mat = db.query(Material).filter(func.upper(Material.code) == raw_code).first()

            if mat:
                # Actualización de ficha existente
                old_stock = float(mat.stock_quantity or 0.0)
                mat.name = raw_name
                mat.category = category
                mat.unit_measure = unit
                mat.min_stock_alert = min_stock
                mat.unit_cost_usd = unit_cost
                mat.location = location
                mat.is_active = True

                # Si el stock cambió (por ejemplo el usuario restó o ajustó)
                if abs(stock_qty - old_stock) > 1e-4:
                    diff = stock_qty - old_stock
                    mov = MaterialMovement(
                        material_id=mat.id,
                        movement_type="ajuste_inventario",
                        quantity=abs(diff),
                        unit_cost_usd=unit_cost,
                        total_cost_usd=round(abs(diff) * unit_cost, 2),
                        destination="Reforma de Inventario (Excel)",
                        notes=f"Ajuste masivo vía Excel: stock anterior {old_stock:.2f} -> nuevo {stock_qty:.2f} ({'+' if diff > 0 else ''}{diff:.2f})",
                        performed_by="Auditoría / Importación Excel"
                    )
                    db.add(mov)
                    mat.stock_quantity = stock_qty
                    mat.total_cost_usd = round(stock_qty * unit_cost, 2)
                    adjusted_stock_count += 1
                else:
                    mat.total_cost_usd = round(mat.stock_quantity * unit_cost, 2)

                updated_count += 1
            else:
                # Creación de nuevo material
                new_mat = Material(
                    code=raw_code,
                    name=raw_name,
                    category=category,
                    unit_measure=unit,
                    stock_quantity=stock_qty,
                    min_stock_alert=min_stock,
                    unit_cost_usd=unit_cost,
                    total_cost_usd=round(stock_qty * unit_cost, 2),
                    location=location,
                    is_active=True
                )
                db.add(new_mat)
                db.flush()

                if stock_qty > 0:
                    mov = MaterialMovement(
                        material_id=new_mat.id,
                        movement_type="ajuste_inventario",
                        quantity=stock_qty,
                        unit_cost_usd=unit_cost,
                        total_cost_usd=round(stock_qty * unit_cost, 2),
                        destination="Carga Inicial Masiva (Excel)",
                        notes=f"Carga inicial de inventario vía Excel ({stock_qty:.2f} {unit})",
                        performed_by="Auditoría / Importación Excel"
                    )
                    db.add(mov)

                created_count += 1

        db.commit()

        return {
            "success": True,
            "created": created_count,
            "updated": updated_count,
            "adjusted_stock": adjusted_stock_count,
            "errors": errors,
            "message": f"Procesamiento completado: {created_count} materiales creados, {updated_count} actualizados ({adjusted_stock_count} con ajuste de stock)."
        }

    # =========================================================================
    # 3. ACTIVOS, HERRAMIENTAS & EQUIPOS: PLANTILLA & EXPORTACIÓN
    # =========================================================================
    @staticmethod
    def generate_assets_template(db: Session) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Activos_Herramientas"
        ws.views.sheetView[0].showGridLines = True

        navy_fill, blue_fill, title_font, header_font, thin_border = InventoryExcelService._create_header_style()
        regular_font = Font(name="Calibri", size=10)
        italic_hint = Font(name="Calibri", size=9, italic=True, color="64748B")

        # Título
        ws.merge_cells("A1:L1")
        ws["A1"] = "DALOR C.A. | SIGO-P - REFORMA Y CATÁLOGO DE HERRAMIENTAS, MAQUINARIAS Y EQUIPOS"
        ws["A1"].font = title_font
        ws["A1"].fill = navy_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:L2")
        ws["A2"] = "Instrucciones: Puede actualizar activos por su CÓDIGO o agregar nuevos. Si agrega una herramienta con CANTIDAD_UNIDADES > 1, el sistema creará automáticamente las unidades correlativas."
        ws["A2"].font = italic_hint
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        headers = [
            ("CODIGO_ACTIVO", 16),
            ("NOMBRE_ACTIVO", 36),
            ("TIPO_ACTIVO", 20),
            ("CATEGORIA", 22),
            ("MARCA", 18),
            ("MODELO", 18),
            ("SERIAL_NUMERO", 20),
            ("PLACA", 14),
            ("UBICACION", 26),
            ("CUSTODIO", 24),
            ("ESTATUS", 18),
            ("CANTIDAD_UNIDADES", 18)
        ]

        ws.row_dimensions[3].height = 24
        for col_idx, (col_name, col_width) in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = blue_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        assets = db.query(Asset).filter(Asset.is_active == True).order_by(Asset.asset_type.asc(), Asset.asset_code.asc()).all()

        current_row = 4
        for a in assets:
            ws.row_dimensions[current_row].height = 19
            row_data = [
                a.asset_code,
                a.name,
                a.asset_type or "herramienta",
                a.category or "General",
                a.brand or "",
                a.model or "",
                a.serial_number or "",
                a.license_plate or "",
                a.current_location or "Sede Central Dalor (Guacara)",
                a.current_custodian_name or "Disponible en Base",
                a.status or "disponible_base",
                1
            ]
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in (1, 3, 8, 11):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 12:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0'
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # =========================================================================
    # 4. ACTIVOS, HERRAMIENTAS & EQUIPOS: IMPORTACIÓN MASIVA
    # =========================================================================
    @staticmethod
    def import_assets_excel(file_bytes: bytes, db: Session) -> Dict[str, Any]:
        wb = load_workbook(filename=io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        created_count = 0
        updated_count = 0
        errors = []

        header_row = 3
        headers_map = {}
        for row_idx in range(1, 6):
            row_vals = [str(ws.cell(row=row_idx, column=c).value or "").strip().upper() for c in range(1, 16)]
            if any("CODIGO" in v for v in row_vals) and any("NOMBRE" in v for v in row_vals):
                header_row = row_idx
                for col_idx in range(1, 16):
                    v = str(ws.cell(row=row_idx, column=col_idx).value or "").strip().upper()
                    if v:
                        headers_map[v] = col_idx
                break

        if not headers_map:
            raise ValueError("El archivo no tiene el formato esperado para activos/herramientas.")

        import re
        max_row = ws.max_row
        for r in range(header_row + 1, max_row + 1):
            def get_val(key: str) -> Optional[Any]:
                col = headers_map.get(key)
                if not col:
                    return None
                val = ws.cell(row=r, column=col).value
                return val

            raw_code = str(get_val("CODIGO_ACTIVO") or get_val("CODIGO") or "").strip().upper()
            raw_name = str(get_val("NOMBRE_ACTIVO") or get_val("NOMBRE") or "").strip()

            if not raw_code and not raw_name:
                continue

            if not raw_name:
                errors.append(f"Fila {r}: Activo {raw_code} no tiene nombre.")
                continue

            a_type = str(get_val("TIPO_ACTIVO") or get_val("TIPO") or "herramienta").strip().lower()
            category = str(get_val("CATEGORIA") or "General").strip()
            brand = str(get_val("MARCA") or "").strip()
            model = str(get_val("MODELO") or "").strip()
            serial = str(get_val("SERIAL_NUMERO") or get_val("SERIAL") or "").strip()
            plate = str(get_val("PLACA") or "").strip().upper()
            loc = str(get_val("UBICACION") or "Sede Central Dalor (Guacara)").strip()
            custodian = str(get_val("CUSTODIO") or "Disponible en Base").strip()
            status = str(get_val("ESTATUS") or "disponible_base").strip().lower()

            try:
                qty = max(1, int(get_val("CANTIDAD_UNIDADES") or get_val("CANTIDAD") or 1))
            except (ValueError, TypeError):
                qty = 1

            # Si viene código existente, actualizar
            asset = db.query(Asset).filter(func.upper(Asset.asset_code) == raw_code).first() if raw_code else None

            if asset:
                asset.name = raw_name
                asset.asset_type = a_type
                asset.category = category
                if brand: asset.brand = brand
                if model: asset.model = model
                if serial: asset.serial_number = serial
                if plate: asset.license_plate = plate
                asset.current_location = loc
                asset.current_custodian_name = custodian
                asset.status = status
                asset.is_active = True
                updated_count += 1
            else:
                # Código nuevo o a generar correlativo
                prefix = raw_code if raw_code else ("HER" if "herramienta" in a_type else "MAQ" if "maq" in a_type else "ACT")
                # Si el usuario mandó N unidades para una herramienta/equipo nuevo
                for unit_i in range(qty):
                    if raw_code and qty == 1:
                        target_code = raw_code
                    else:
                        base_pref = re.sub(r'[-_]?\d+$', '', prefix).rstrip('-_') or "HER"
                        existing = set(c[0] for c in db.query(Asset.asset_code).filter(Asset.asset_code.like(f"{base_pref}%")).all())
                        c_idx = 1
                        candidate = f"{base_pref}-{c_idx:03d}"
                        while candidate in existing:
                            c_idx += 1
                            candidate = f"{base_pref}-{c_idx:03d}"
                        target_code = candidate

                    new_asset = Asset(
                        asset_code=target_code,
                        name=raw_name,
                        asset_type=a_type,
                        category=category,
                        brand=brand,
                        model=model,
                        serial_number=serial if qty == 1 else (f"{serial}-{unit_i+1}" if serial else f"S/N-{target_code}"),
                        license_plate=plate or None,
                        current_location=loc,
                        current_custodian_name=custodian,
                        status=status,
                        is_exclusive=True,
                        is_active=True
                    )
                    db.add(new_asset)
                    db.flush()
                    created_count += 1

        db.commit()

        return {
            "success": True,
            "created": created_count,
            "updated": updated_count,
            "errors": errors,
            "message": f"Procesamiento completado: {created_count} activos/herramientas creados, {updated_count} actualizados."
        }
